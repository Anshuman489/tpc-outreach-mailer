"""Excel helpers: reading recipients, saving status, red row highlighting."""
import re

import openpyxl
import pandas as pd
from openpyxl.styles import PatternFill

RED_FILL = PatternFill(start_color="FF0000", end_color="FF0000", fill_type="solid")
RED_COLUMNS = range(1, 6)  # A=1 to E=5


def read_sheet(path):
    """Read a sheet with headers normalised to stripped lowercase."""
    df = pd.read_excel(path)
    df.columns = df.columns.str.strip().str.lower()
    return df


def save_sheet(df, path):
    df.to_excel(path, index=False)


def split_recipients(raw):
    """Split a cell holding several addresses separated by / ; , or whitespace."""
    return [e.strip() for e in re.split(r'[\/;,\s]+', raw) if e.strip()]


def recipient_set(path):
    """All lowercase addresses in the sheet's To column."""
    df = read_sheet(path)
    sheet_emails = set()
    for val in df["to"].dropna():
        for email in split_recipients(str(val)):
            sheet_emails.add(email.lower())
    print(f"Loaded {len(sheet_emails)} emails from sheet.")
    return sheet_emails


def open_workbook(path):
    wb = openpyxl.load_workbook(path)
    return wb, wb.active


def fill_row_red(ws, row_num, skip_col=None):
    for col in RED_COLUMNS:
        if col != skip_col:
            ws.cell(row=row_num, column=col).fill = RED_FILL
